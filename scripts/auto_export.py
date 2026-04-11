#!/usr/bin/env python3
"""
Auto Export Script - Daily Backup at 10:00 PM

Automatically exports all TaxCore data to Excel files daily.
Saves to /app/exports/ with timestamp.
"""

import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import pandas as pd
    from pymongo import MongoClient
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils.dataframe import dataframe_to_rows
except ImportError as e:
    print(f"❌ Missing required package: {e}")
    print("Installing required packages...")
    os.system("pip install pandas openpyxl pymongo")
    print("Please run the script again.")
    sys.exit(1)

# MongoDB Connection
MONGO_URL = os.getenv("MONGO_URL", "mongodb://localhost:27017")
EXPORT_DIR = "/app/exports"

def ensure_export_directory():
    """Create exports directory if it doesn't exist"""
    Path(EXPORT_DIR).mkdir(parents=True, exist_ok=True)
    print(f"✅ Export directory ready: {EXPORT_DIR}")

def get_timestamp():
    """Get formatted timestamp for filename"""
    return datetime.now().strftime("%Y%m%d_%H%M%S")

def get_readable_timestamp():
    """Get human-readable timestamp"""
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

def export_to_excel():
    """Export all MongoDB data to Excel file"""
    try:
        print(f"\n🔄 Starting auto-export at {get_readable_timestamp()}")
        
        # Connect to MongoDB
        client = MongoClient(MONGO_URL)
        db = client.taxcore
        
        # Create Excel workbook
        timestamp = get_timestamp()
        filename = f"TaxCore_AutoExport_{timestamp}.xlsx"
        filepath = os.path.join(EXPORT_DIR, filename)
        
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet
        
        # Define collections to export
        collections = {
            "Clients": "clients",
            "Users": "users",
            "Documents": "documents",
            "Work Processing": "work",
            "Billing": "billing",
            "Firm Accounts": "firm_accounts",
            "Audit Logs": "audit_logs",
        }
        
        total_records = 0
        
        # Export each collection
        for sheet_name, collection_name in collections.items():
            print(f"  📊 Exporting {sheet_name}...")
            
            # Fetch data from MongoDB
            data = list(db[collection_name].find({}, {"_id": 0}))
            
            if not data:
                print(f"    ⚠️  No data in {sheet_name}")
                continue
            
            # Convert to DataFrame
            df = pd.DataFrame(data)
            total_records += len(df)
            
            # Create worksheet
            ws = wb.create_sheet(title=sheet_name)
            
            # Style header row
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=11)
            
            # Write data to worksheet
            for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
                for c_idx, value in enumerate(row, 1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=value)
                    
                    # Style header
                    if r_idx == 1:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            print(f"    ✅ Exported {len(df)} records")
        
        # Add summary sheet
        ws_summary = wb.create_sheet(title="Export Info", index=0)
        ws_summary.append(["TaxCore Data Export"])
        ws_summary.append(["Export Date:", get_readable_timestamp()])
        ws_summary.append(["Total Records:", total_records])
        ws_summary.append([""])
        ws_summary.append(["Collections Exported:"])
        for sheet_name, collection_name in collections.items():
            count = len(list(db[collection_name].find({})))
            ws_summary.append([sheet_name, count])
        
        # Style summary sheet
        ws_summary['A1'].font = Font(bold=True, size=14)
        for row in ws_summary.iter_rows(min_row=2, max_row=3, min_col=1, max_col=1):
            for cell in row:
                cell.font = Font(bold=True)
        
        # Save workbook
        wb.save(filepath)
        
        # Get file size
        file_size = os.path.getsize(filepath) / 1024  # KB
        
        print(f"\n✅ Export completed successfully!")
        print(f"  📁 File: {filename}")
        print(f"  💾 Size: {file_size:.2f} KB")
        print(f"  📊 Records: {total_records}")
        print(f"  📂 Location: {filepath}")
        
        # Close MongoDB connection
        client.close()
        
        # Clean up old exports (keep last 30 days)
        cleanup_old_exports()
        
        return filepath
        
    except Exception as e:
        print(f"❌ Export failed: {str(e)}")
        import traceback
        traceback.print_exc()
        return None

def cleanup_old_exports():
    """Delete export files older than 30 days"""
    try:
        print(f"\n🧹 Cleaning up old exports...")
        
        current_time = datetime.now().timestamp()
        days_to_keep = 30
        cutoff_time = current_time - (days_to_keep * 24 * 60 * 60)
        
        deleted_count = 0
        for filename in os.listdir(EXPORT_DIR):
            if filename.startswith("TaxCore_AutoExport_") and filename.endswith(".xlsx"):
                filepath = os.path.join(EXPORT_DIR, filename)
                file_time = os.path.getmtime(filepath)
                
                if file_time < cutoff_time:
                    os.remove(filepath)
                    deleted_count += 1
                    print(f"  🗑️  Deleted: {filename}")
        
        if deleted_count > 0:
            print(f"✅ Cleaned up {deleted_count} old export(s)")
        else:
            print(f"✅ No old exports to clean up")
        
    except Exception as e:
        print(f"⚠️  Cleanup warning: {str(e)}")

def main():
    """Main execution function"""
    print("=" * 60)
    print("TaxCore Automated Daily Export")
    print("=" * 60)
    
    # Ensure export directory exists
    ensure_export_directory()
    
    # Run export
    export_file = export_to_excel()
    
    if export_file:
        print("\n" + "=" * 60)
        print("✅ Automated export completed successfully!")
        print("=" * 60)
        return 0
    else:
        print("\n" + "=" * 60)
        print("❌ Automated export failed!")
        print("=" * 60)
        return 1

if __name__ == "__main__":
    sys.exit(main())
